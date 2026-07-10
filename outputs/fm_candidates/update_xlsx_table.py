"""Update the existing CPS FM to Trace Gen sheet in cps_fm_candidates.xlsx
with the latest 5-fold test metrics (median [min-max]) and timing-pass p99
ms/tick. Applies native Excel bold to the best variant per column.

Layout (preserves existing headers):
  cols A-E:  ID, Sensor Encoder, Video Encoder, Fusion, Trace Generator
  cols F-I:  NED state-0/1/2/all
  cols J-M:  Exact Match % state-0/1/2/all      (auto × 100 with %)
  cols N-Q:  Vendi-ratio state-0/1/2/all
  col R:     CrystalBLEU (all)
  col S:     Self-BLEU_gen (all)
  col T:     ROUGE-L F1 (all)
  col U:     Token-gen accuracy %  (auto × 100 with %)
  col V:     Inference Time p99 ms/tick H100 (from timing-only JSON)

Variant rows (already present in the sheet):
  row 4: BL1   = MDLM
  row 5: BL2   = AR-modern
  row 6: BL3   = Chronos + V-JEPA 2.1 (frozen)
  row 7: BL3b  = LoRA V-JEPA  (skipped — N/A everywhere)
  row 8: BL4   = Chronos + CoTracker3
  row 9: BL5   = MOMENT-1
"""
from pathlib import Path
import json
from typing import Optional

import numpy as np
import openpyxl
from openpyxl.styles import Font

XLSX = Path("/home/simran/allspark-data-exploration/CPS-Debugger/outputs/fm_candidates/results/cps_fm_candidates.xlsx")
RESULTS = Path("/home/simran/allspark-data-exploration/CPS-Debugger/outputs/fm_candidates/results")

# Row -> (json filename pattern, timing-only json relative path or None for N/A)
ROW_MAP = {
    4: ("mdlm_fold_{fold}/mdlm_eval_test.json",            "timing_mdlm_fold_0/mdlm_eval_test.json"),
    5: ("ar_modern_fold_{fold}/ar_modern_eval_test.json",  "timing_ar_modern_fold_0/ar_modern_eval_test.json"),
    6: ("bl3_fold_{fold}/bl3_eval_test.json",              "timing_bl3_fold_0/bl3_eval_test.json"),
    7: (None,                                              None),
    8: ("bl4_fold_{fold}/bl4_eval_test.json",              "timing_bl4_fold_0/bl4_eval_test.json"),
    9: ("bl6_fold_{fold}/bl6_eval_test.json",              "timing_bl6_fold_0/bl6_eval_test.json"),
}

# (col_letter, metric_key, fmt, "lower"/"higher"/"near_one", percent_x100)
# percent_x100 = True means multiply by 100 and append "%"
COLUMN_SPEC = [
    # NED (F-I): per-state s0/s1/s2/all
    ("F", ("ned_mean", "state", 0), "{:.4f}", "lower",   False),
    ("G", ("ned_mean", "state", 1), "{:.4f}", "lower",   False),
    ("H", ("ned_mean", "state", 2), "{:.4f}", "lower",   False),
    ("I", ("ned_mean", "all",   None), "{:.4f}", "lower",  False),
    # Exact Match % (J-M)
    ("J", ("exact_match", "state", 0), "{:.2f}%", "higher", True),
    ("K", ("exact_match", "state", 1), "{:.2f}%", "higher", True),
    ("L", ("exact_match", "state", 2), "{:.2f}%", "higher", True),
    ("M", ("exact_match", "all",   None), "{:.2f}%", "higher", True),
    # Vendi-ratio (N-Q)
    ("N", ("vendi_ratio_gen_over_ref", "state", 0), "{:.3f}", "near_one", False),
    ("O", ("vendi_ratio_gen_over_ref", "state", 1), "{:.3f}", "near_one", False),
    ("P", ("vendi_ratio_gen_over_ref", "state", 2), "{:.3f}", "near_one", False),
    ("Q", ("vendi_ratio_gen_over_ref", "all",   None), "{:.3f}", "near_one", False),
    # Single-column metrics
    ("R", ("crystal_bleu",         "all", None), "{:.4f}",   "higher", False),
    ("S", ("self_bleu_generated",  "all", None), "{:.4f}",   "lower",  False),
    ("T", ("rouge_l_f1_mean",      "all", None), "{:.4f}",   "higher", False),
    ("U", ("token_acc_generation", "all", None), "{:.2f}%",  "higher", True),
    # Timing — single fold from timing JSON
    ("V", ("p99_ms_per_tick",      "timing", None), "{:.3f}", "lower", False),
]


def _scalar(j, key, mode_prefixes=("sampled_", "baseline_", "greedy_", "")) -> Optional[float]:
    if j is None:
        return None
    for p in mode_prefixes:
        k = f"{p}{key}"
        if k in j and isinstance(j[k], (int, float)):
            return float(j[k])
    return None


def _per_state(j, key, state, mode_prefixes=("sampled_", "baseline_", "greedy_", "")) -> Optional[float]:
    if j is None:
        return None
    for p in mode_prefixes:
        ps_key = f"{p}per_state"
        if ps_key in j and isinstance(j[ps_key], dict):
            sd = j[ps_key].get(str(state))
            if sd and key in sd and isinstance(sd[key], (int, float)):
                return float(sd[key])
    return None


def _try_load(path: Path):
    if not path.exists():
        return None
    try:
        return json.loads(path.read_text())
    except Exception:
        return None


def _agg_cell(values, fmt, percent_x100) -> str:
    arr = np.array([v for v in values if v is not None], dtype=float)
    if len(arr) == 0:
        return ""
    if percent_x100:
        arr = arr * 100.0
    med = float(np.median(arr))
    lo, hi = float(arr.min()), float(arr.max())
    return f"{fmt.format(med)} [{fmt.format(lo)}-{fmt.format(hi)}]"


def _single_cell(value, fmt, percent_x100) -> str:
    if value is None:
        return ""
    v = value * 100.0 if percent_x100 else value
    return fmt.format(v)


def _load_metric_for_row(pattern, timing_path, metric_spec):
    """Return list of fold values for this metric (length 5; or 1 for timing)."""
    key, kind, state = metric_spec
    if kind == "timing":
        if timing_path is None:
            return []
        tj = _try_load(RESULTS / timing_path)
        v = _scalar(tj, key)
        return [v] if v is not None else []
    # per-fold metric
    if pattern is None:
        return []
    fold_jsons = [_try_load(RESULTS / pattern.format(fold=f)) for f in range(5)]
    if kind == "all":
        return [_scalar(j, key) for j in fold_jsons]
    if kind == "state":
        return [_per_state(j, key, state) for j in fold_jsons]
    return []


def _parse_median(cell_text: str) -> Optional[float]:
    if not cell_text or cell_text == "N/A":
        return None
    s = cell_text.replace("%", "").strip()
    if "[" in s:
        s = s.split("[", 1)[0].strip()
    try:
        return float(s)
    except ValueError:
        return None


def _best_index(values_text, direction):
    nums = [(_parse_median(t), i) for i, t in enumerate(values_text)]
    nums = [(n, i) for n, i in nums if n is not None]
    if not nums:
        return None
    if direction == "lower":
        return min(nums)[1]
    if direction == "higher":
        return max(nums, key=lambda t: t[0])[1]
    if direction == "near_one":
        return min(nums, key=lambda t: abs(t[0] - 1.0))[1]
    return None


def main():
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["CPS FM to Trace Gen"]

    # Update the U-column header (row 2) to include % since we're flipping units
    if ws.cell(row=2, column=21).value and "Token" in str(ws.cell(row=2, column=21).value):
        cur = str(ws.cell(row=2, column=21).value).rstrip()
        if "%" not in cur:
            ws.cell(row=2, column=21).value = cur + " %"

    # Build the cell-text matrix per (row, col)
    # rows is row-index in Excel (4..9), cells dict: col_letter -> text
    rows_data = {}
    for excel_row, (pattern, timing_path) in ROW_MAP.items():
        cells = {}
        if pattern is None:
            # N/A row (BL3b)
            for col_letter, _, _, _, _ in COLUMN_SPEC:
                cells[col_letter] = "N/A"
        else:
            for col_letter, mspec, fmt, _, percent_x100 in COLUMN_SPEC:
                key, kind, _state = mspec
                vals = _load_metric_for_row(pattern, timing_path, mspec)
                if kind == "timing":
                    cells[col_letter] = _single_cell(vals[0] if vals else None, fmt, percent_x100)
                else:
                    cells[col_letter] = _agg_cell(vals, fmt, percent_x100)
        rows_data[excel_row] = cells

    # Write text into cells; clear bold first (start clean)
    for excel_row, cells in rows_data.items():
        for col_letter, text in cells.items():
            cell = ws[f"{col_letter}{excel_row}"]
            cell.value = text
            # Reset font (no bold) — we'll reapply on bests next
            f = cell.font
            cell.font = Font(name=f.name, size=f.size, bold=False, italic=f.italic,
                             color=f.color, vertAlign=f.vertAlign, underline=f.underline,
                             strike=f.strike)

    # Apply bold to best per column
    for col_letter, _, _, direction, _ in COLUMN_SPEC:
        col_values = [rows_data[r][col_letter] for r in sorted(rows_data)]
        best_local = _best_index(col_values, direction)
        if best_local is None:
            continue
        excel_row = sorted(rows_data)[best_local]
        cell = ws[f"{col_letter}{excel_row}"]
        f = cell.font
        cell.font = Font(name=f.name, size=f.size, bold=True, italic=f.italic,
                         color=f.color, vertAlign=f.vertAlign, underline=f.underline,
                         strike=f.strike)

    # Bold all the headers (rows 1-3)
    for r in (1, 2, 3):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            if cell.value is None:
                continue
            f = cell.font
            cell.font = Font(name=f.name, size=f.size, bold=True, italic=f.italic,
                             color=f.color, vertAlign=f.vertAlign, underline=f.underline,
                             strike=f.strike)

    wb.save(XLSX)
    print(f"Saved {XLSX}")
    # Quick summary of what landed
    print("\nUpdated rows:")
    for r in sorted(rows_data):
        ident = ws.cell(row=r, column=1).value
        sample = ws.cell(row=r, column=ord("I") - ord("A") + 1).value  # NED-all
        print(f"  row {r} ({ident}): NED-all = {sample!r}")


if __name__ == "__main__":
    main()
